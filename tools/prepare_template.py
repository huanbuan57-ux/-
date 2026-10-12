"""从旧版技术通知创建 DOCX 自动化模板，原件始终只读。

用法：python tools/prepare_template.py --source 技术通知.doc --output data/自动化模板.docx
旧版 .doc 需要 LibreOffice：可通过 --soffice 指定 soffice.exe。
也可直接把已经真实转换的 .docx 作为 --source。
模板保留首页通知表、签批、发往部门、原始包装文字和页眉水印，
移除旧两产品明细及空白页，改成每产品一行的双栏信息块。
"""
from __future__ import annotations

import argparse
import hashlib
from pathlib import Path
import shutil
import subprocess
import tempfile

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def convert_source(source: Path, destination: Path, soffice: str | None) -> Path:
    if source.suffix.lower() == ".docx":
        shutil.copy2(source, destination / source.name)
        return destination / source.name
    binary = soffice or shutil.which("soffice.exe") or shutil.which("soffice")
    if not binary:
        raise RuntimeError("转换旧版 .doc 需要 LibreOffice，请用 --soffice 指定可执行文件")
    profile = (destination / "lo-profile").resolve().as_uri()
    subprocess.run([binary, f"-env:UserInstallation={profile}", "--headless", "--convert-to",
                    "docx", "--outdir", str(destination), str(source)], check=True,
                   stdout=subprocess.PIPE, stderr=subprocess.PIPE)
    converted = destination / f"{source.stem}.docx"
    if not converted.exists():
        raise RuntimeError("LibreOffice 未生成转换文件")
    return converted


def replace_cell(cell, text: str) -> None:
    paragraph = cell.paragraphs[0]
    if paragraph.runs:
        paragraph.runs[0].text = text
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(text)
    for extra in cell.paragraphs[1:]:
        extra._element.getparent().remove(extra._element)


def build_template(converted: Path, output: Path) -> None:
    doc = Document(converted)
    if not doc.tables:
        raise ValueError("原通知缺少首页表格")
    front = doc.tables[0]
    # Source has a 33-column grid with merged cells. Address stable unique
    # positions without rebuilding the grid or its signatures/departments.
    replace_cell(front.cell(0, 21), "{{通知编号}}")
    replace_cell(front.cell(2, 0), "产品明细见次页")
    replace_cell(front.cell(3, 26), "{{实施日期}}")

    # The converted source contains dozens of blank paragraphs and a second
    # page break. Remove only body elements outside the preserved first table.
    body = doc.element.body
    for element in list(body):
        if element is not front._element and element.tag != qn("w:sectPr"):
            body.remove(element)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.space_before = Pt(0)
    p.add_run().add_break(WD_BREAK.PAGE)
    p = doc.add_paragraph("出口散件产品明细")
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.keep_with_next = True
    p.runs[0].font.name = "宋体"
    p.runs[0].font.size = Pt(14)
    p.runs[0].bold = True
    p = doc.add_paragraph("合同号：{{合同号}}    车型：{{车型}}")
    p.paragraph_format.space_after = Pt(8)
    p = doc.add_paragraph("交货地点：{{delivery_place}}    交货时间：{{delivery_time}}")
    p.paragraph_format.space_after = Pt(8)

    table = doc.add_table(rows=1, cols=2)
    table.autofit = False
    table.columns[0].width = Cm(14)
    table.columns[1].width = Cm(14)
    table.cell(0, 0).width = Cm(14)
    table.cell(0, 1).width = Cm(14)
    values = [
        "图号：{{图号}}\n中文名称：{{中文名称}}\n英文名称：{{英文名称}}\n数量：{{数量}}\n包装模数：{{包装模数}} 台套一组\n单箱数量：{{单箱数量}}\n纸箱总数：{{纸箱总数}}",
        "工位号：{{工位号}}\n纸箱尺寸：{{纸箱尺寸}}\n重量说明：{{重量说明}}\nLOGO类型：{{LOGO类型}}\n国家：{{国家}}\n合同号：{{合同号}}\n车型：{{车型}}",
    ]
    for cell, value in zip(table.rows[0].cells, values):
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        cell.text = value
        tc_pr = cell._tc.get_or_add_tcPr()
        borders = OxmlElement("w:tcBorders")
        for edge in ("top", "left", "bottom", "right"):
            e = OxmlElement(f"w:{edge}")
            e.set(qn("w:val"), "single")
            e.set(qn("w:sz"), "4")
            e.set(qn("w:color"), "888888")
            borders.append(e)
        tc_pr.append(borders)
        margins = OxmlElement("w:tcMar")
        for edge in ("top", "left", "bottom", "right"):
            e = OxmlElement(f"w:{edge}")
            e.set(qn("w:w"), "100")
            e.set(qn("w:type"), "dxa")
            margins.append(e)
        tc_pr.append(margins)
        for p in cell.paragraphs:
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.05
            for run in p.runs:
                run.font.name = "宋体"
                run._element.get_or_add_rPr().append(OxmlElement("w:rFonts"))
                run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
                run.font.size = Pt(11)
    cant_split = OxmlElement("w:cantSplit")
    table.rows[0]._tr.get_or_add_trPr().append(cant_split)
    p = doc.add_paragraph("通用技术要求：{{通用技术要求}}")
    p.paragraph_format.space_before = Pt(8)
    doc.add_paragraph("完成日期：{{finish_date}}")
    doc.add_paragraph("标签网址：{{标签网址}}")
    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--source", type=Path, default=Path.home() / "Desktop" / "技术通知.doc")
    parser.add_argument("--output", type=Path, default=Path("data") / "自动化模板.docx")
    parser.add_argument("--soffice")
    args = parser.parse_args()
    before = sha256(args.source)
    with tempfile.TemporaryDirectory(prefix="notice_template_") as td:
        converted = convert_source(args.source.resolve(), Path(td), args.soffice)
        build_template(converted, args.output)
    if sha256(args.source) != before:
        raise RuntimeError("原文件校验失败")
    print(f"模板已生成：{args.output.resolve()}；原文件 SHA-256 保持不变：{before}")


if __name__ == "__main__":
    main()
