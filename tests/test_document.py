import tempfile
import unittest
from pathlib import Path

from docx import Document

from export_notice.document import render_notification, validate_template


REQUIRED = [
    "{{number}}", "{{date}}", "{{contract_no}}", "{{model}}",
    "{{general_requirements}}", "{{label_url}}", "{{row.code}}",
    "{{row.cn}}", "{{row.en}}", "{{row.quantity}}", "{{row.modulus}}",
    "{{row.units}}", "{{row.cartons}}", "{{row.station}}", "{{row.size}}",
    "{{row.weight}}", "{{row.logo_type}}", "{{row.country}}",
    "{{row.contract_no}}", "{{row.model}}",
]


def make_template(path: Path, include_all=True):
    doc = Document()
    doc.add_paragraph("通知 {{number}} {{date}} {{contract_no}} {{model}}")
    doc.add_paragraph("要求 {{general_requirements}} 标签 {{label_url}}")
    table = doc.add_table(rows=1, cols=2)
    row = table.rows[0]
    row.cells[0].text = "图号 {{row.code}} / {{row.cn}} / {{row.en}} / {{row.quantity}} / {{row.modulus}} / {{row.units}} / {{row.cartons}}"
    row.cells[1].text = "工位 {{row.station}} 尺寸 {{row.size}} 重量 {{row.weight}} LOGO {{row.logo_type}} 国家 {{row.country}} 合同 {{row.contract_no}} 车型 {{row.model}}"
    if not include_all:
        row.cells[0].text = "{{row.code}}"
    doc.save(path)


class DocumentTests(unittest.TestCase):
    def test_actual_template_contains_all_slots_and_no_source_batch_values(self):
        template = Path(__file__).resolve().parents[1] / "data" / "自动化模板.docx"
        self.assertEqual(validate_template(template), [])
        doc = Document(template)
        texts = [p.text for p in doc.paragraphs]
        texts += [c.text for t in doc.tables for r in t.rows for c in r.cells]
        text = "\n".join(texts)
        for stale in ("S111900000182", "S111900000294", "HWCSC26278YX", "C310026S001YX00CBP", "2026.10.10", "以上2026/10/16"):
            self.assertNotIn(stale, text)
        for label in ("编制", "审核", "批准", "生产计划部", "质量部"):
            self.assertIn(label, text)

    def test_split_run_replacement_preserves_surrounding_font_and_header(self):
        with tempfile.TemporaryDirectory() as td:
            template = Path(td) / "template.docx"
            make_template(template)
            doc = Document(template)
            p = doc.paragraphs[0]
            p.clear()
            p.add_run("加粗原文 ").bold = True
            p.add_run("{{num").italic = True
            p.add_run("ber}}")
            p.add_run(" 保留尾文 {{date}} {{contract_no}} {{model}}").underline = True
            doc.sections[0].header.paragraphs[0].text = "页眉 {{number}}"
            doc.save(template)
            data = {"number": "TN-9", "date": "2026", "contract_no": "C", "model": "M",
                    "general_requirements": "保持", "label_url": "https://x", "rows": []}
            output = Path(td) / "out.docx"
            render_notification(template, data, output)
            doc = Document(output)
            self.assertTrue(doc.paragraphs[0].runs[0].bold)
            self.assertTrue(doc.paragraphs[0].runs[-1].underline)
            self.assertIn("TN-9", doc.sections[0].header.paragraphs[0].text)

    def test_chinese_slots_support_nested_globals_and_row_dicts(self):
        with tempfile.TemporaryDirectory() as td:
            template = Path(td) / "template.docx"
            doc = Document()
            doc.add_paragraph("{{通知编号}} {{实施日期}} {{合同号}} {{车型}} {{通用技术要求}} {{标签网址}}")
            table = doc.add_table(rows=1, cols=2)
            table.cell(0, 0).text = "{{图号}} {{中文名称}} {{英文名称}} {{数量}} {{包装模数}} {{单箱数量}} {{纸箱总数}}"
            table.cell(0, 1).text = "{{工位号}} {{纸箱尺寸}} {{重量说明}} {{LOGO类型}} {{国家}} {{合同号}} {{车型}}"
            doc.save(template)
            self.assertEqual(validate_template(template), [])
            row = {"图号": "P3", "中文名称": "产品", "英文名称": "Product", "数量": 2,
                   "包装模数": 1, "单箱数量": 2, "纸箱总数": 1, "工位号": "/", "纸箱尺寸": "1x2",
                   "重量说明": "1kg", "LOGO类型": "L", "国家": "中国", "合同号": "C", "车型": "M"}
            output = Path(td) / "out.docx"
            render_notification(template, {"globals": {"通知编号": "N", "实施日期": "D", "合同号": "C",
                                "车型": "M", "通用技术要求": "原文", "标签网址": "http://x"}, "rows": [row]}, output)
            self.assertIn("P3", Document(output).tables[0].cell(0, 0).text)

    def test_unresolved_values_do_not_write_output(self):
        with tempfile.TemporaryDirectory() as td:
            template = Path(td) / "template.docx"
            make_template(template)
            output = Path(td) / "out.docx"
            with self.assertRaisesRegex(ValueError, "未解析"):
                render_notification(template, {"rows": []}, output)
            self.assertFalse(output.exists())

    def test_validate_template_reports_missing_placeholders(self):
        with tempfile.TemporaryDirectory() as td:
            path = Path(td) / "bad.docx"
            make_template(path, include_all=False)
            missing = validate_template(path)
            self.assertIn("{{row.cn}}", missing)

    def test_render_repeats_detail_row_and_replaces_globals(self):
        with tempfile.TemporaryDirectory() as td:
            template = Path(td) / "template.docx"
            output = Path(td) / "output.docx"
            make_template(template)
            data = {
                "number": "TN-001", "date": "2026-10-12", "contract_no": "C-1",
                "model": "M1", "country": "中国", "logo_type": "标准",
                "general_requirements": "保持原包装要求", "label_url": "https://x",
                "rows": [
                    {"code": "P1", "cn": "零件一", "en": "Part One", "quantity": 2,
                     "modulus": 1, "units": 2, "cartons": 1, "station": "A", "size": "1x2",
                     "weight": "1kg", "logo_type": "A", "country": "中国", "contract_no": "C-1", "model": "M1"},
                    {"code": "P2", "cn": "零件二", "en": "Part Two", "quantity": 4,
                     "modulus": 2, "units": 2, "cartons": 2, "station": "B", "size": "2x3",
                     "weight": "2kg", "logo_type": "B", "country": "德国", "contract_no": "C-2", "model": "M2"},
                ],
            }
            result = render_notification(template, data, output)
            self.assertEqual(result, output)
            rendered = Document(output)
            text = "\n".join(p.text for p in rendered.paragraphs)
            text += "\n" + "\n".join(c.text for t in rendered.tables for r in t.rows for c in r.cells)
            self.assertIn("TN-001", text)
            self.assertIn("P1", text)
            self.assertIn("P2", text)
            self.assertNotIn("{{", text)

    def test_render_rejects_missing_template_slot_and_unresolved_data(self):
        with tempfile.TemporaryDirectory() as td:
            template = Path(td) / "template.docx"
            make_template(template, include_all=False)
            with self.assertRaises(ValueError):
                render_notification(template, {"rows": []}, Path(td) / "out.docx")


if __name__ == "__main__":
    unittest.main()
