import tempfile
import unittest
from pathlib import Path
from openpyxl import Workbook
from export_notice.service import analyze_batch, analyze_rows, notification_data
from export_notice.core import BatchRow, ProductMaster, ProductRecord
from export_notice.io import load_batch

ROOT = Path(__file__).resolve().parents[1]
class RegressionTests(unittest.TestCase):
    def test_actual_batch_matches_completed_notice(self):
        result=analyze_batch(ROOT/'examples'/'样例批次',ROOT/'data'/'产品资料表.xlsx')
        self.assertEqual(result.errors,[])
        self.assertEqual(len(result.all_rows),5)
        self.assertEqual(len(result.excluded),3)
        self.assertEqual({r.standard_code:r.carton_count for r in result.rows},
                         {'S111900000182':60,'S111900000294':5})
        self.assertEqual(result.requirements.country,'/')
        self.assertEqual(result.requirements.logo_type,'/')
        self.assertEqual(notification_data(result,'N','2026-10-12')['finish_date'],'2026-10-16')
    def test_master_modulus_fallback_is_valid(self):
        master=ProductMaster({'A':ProductRecord('A',chinese_name='A',english_name='Part',
               packaging_modulus=36,units_per_carton=2,carton_size='1*2*3',weight_note='实际')})
        result=analyze_rows([BatchRow('A','A',12,{},'x.xlsx','表',2)],master,
            ['合同号：ABC 车型：C1 国家：/ LOGO类型：/ 标签网址：https://example.com'])
        self.assertEqual(result.errors,[])
        self.assertEqual(result.rows[0].resolved['包装模数'],'36')
    def test_xls_multiple_sheets_are_read_independently(self):
        import xlwt
        with tempfile.TemporaryDirectory() as td:
            wb=xlwt.Workbook()
            for name,code,qty in [('一','A',12),('二','B',24)]:
                s=wb.add_sheet(name)
                for c,v in enumerate(['图号','数量']): s.write(0,c,v)
                s.write(1,0,code); s.write(1,1,qty)
            wb.save(str(Path(td)/'a.xls'))
            rows=load_batch(td).rows
            self.assertEqual([(r.raw_code,r.quantity) for r in rows],[('A',12.0),('B',24.0)])
    def test_remark_fields_and_repeated_header(self):
        with tempfile.TemporaryDirectory() as td:
            w=Workbook();s=w.active
            s.append(['图号','数量','备注'])
            s.append(['A',12,'英文名称：Part A；每箱数量：2；纸箱尺寸：100*200*300；工位号：W1；包装模数：12'])
            s.append(['图号','数量','备注'])
            s.append(['B',24,'工位号：W2'])
            w.save(Path(td)/'a.xlsx')
            rows=load_batch(td).rows
            self.assertEqual(len(rows),2)
            self.assertEqual(rows[0].fields['英文名称'],'Part A')
            self.assertEqual(rows[0].fields['工位号'],'W1')
            self.assertEqual(rows[0].fields['单箱数量'],'2')

    def test_structured_columns_and_sources(self):
        with tempfile.TemporaryDirectory() as td:
            w=Workbook();s=w.active
            s.append(['图号','数量','合同号','车型','国家','LOGO类型','英文名称','包装模数','单箱数量','纸箱尺寸'])
            s.append(['A',12,'ABC','C1','/','/', 'Part',24,2,'1*2*3'])
            w.save(Path(td)/'a.xlsx')
            batch=load_batch(td)
            self.assertIn('合同号：ABC',batch.requirement_texts)
            self.assertEqual(batch.requirement_sources[0]['row'],2)
    def test_missing_template_writes_report_only(self):
        from export_notice.service import generate
        result=analyze_batch(ROOT/'examples'/'样例批次',ROOT/'data'/'产品资料表.xlsx')
        with tempfile.TemporaryDirectory() as td:
            doc,report,errors=generate(result,Path(td)/'missing.docx',td,'batch')
            self.assertIsNone(doc)
            self.assertTrue(report.exists())
            self.assertTrue(errors)
    def test_excel_packaging_override_and_station_priority(self):
        master=ProductMaster({'A':ProductRecord('A',chinese_name='A',english_name='Master',
             packaging_modulus=12,units_per_carton=2,carton_size='old',station_no='old',weight_note='实际')})
        row=BatchRow('A+ck','A',12,{'英文名称':'Excel English','包装模数':24,'单箱数量':3,'纸箱尺寸':'new','工位号':'W7'},'a','s',1)
        result=analyze_rows([row],master,['合同号：ABC 车型：C1 国家：/ LOGO类型：/ 标签网址：https://example.com'])
        self.assertEqual(result.errors,[])
        self.assertEqual(row.carton_count,4)
        self.assertEqual(row.resolved['工位号'],'W7')
        self.assertEqual(row.resolved['英文名称'],'Excel English')
        self.assertEqual(len(result.warnings),4)
