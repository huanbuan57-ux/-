import tempfile
import unittest
from pathlib import Path
from openpyxl import Workbook
from export_notice.io import load_batch, load_product_master
from export_notice.requirements import extract_requirements

class InputTests(unittest.TestCase):
    def test_multisheet_aliases_and_unsupported_file_block(self):
        with tempfile.TemporaryDirectory() as td:
            w = Workbook()
            s = w.active
            s.append(['图号','数量','产品英文名称','每箱数量','经济批量','箱尺寸','工位'])
            s.append(['A+ck',24,'Hello',2,12,'100*200*300','W1'])
            s = w.create_sheet('另一表')
            s.append(['物料编码','计划数量'])
            s.append(['B-ck',36])
            w.save(Path(td)/'a.xlsx')
            batch = load_batch(td)
            self.assertEqual(len(batch.rows),2)
            self.assertEqual(batch.rows[0].fields['英文名称'],'Hello')
            self.assertEqual(batch.rows[0].fields['工位号'],'W1')
            (Path(td)/'bad.xls').write_bytes(b'broken')
            self.assertTrue(load_batch(td).errors)

    def test_explicit_label_prohibition_overrides_destination(self):
        req = extract_requirements(['KD散件/HWCSC26278YX-1/土耳其-TR/C310026S001YX00CBP//出口指定',
          '包装标签上不得体现福田Logo及国家信息。logo类型：无（选第二个） 国家：/；包装模数：24；标签网址：https://example.com/a'])
        self.assertEqual(req.contract_no,'HWCSC26278YX')
        self.assertEqual(req.country,'/')
        self.assertEqual(req.logo_type,'/')
        self.assertEqual(req.modulus,'24')
        self.assertEqual(req.label_url,'https://example.com/a')

    def test_unknown_country_and_conflicting_contract_are_blocked(self):
        self.assertTrue(extract_requirements(['合同号：ABC 车型：C1 包装模数：24']).errors)
        req = extract_requirements(['合同号：ABC 车型：C1 国家：土耳其 LOGO类型：福田',
                                    '合同号：DEF'])
        self.assertTrue(any('合同号' in x for x in req.errors))

    def test_invalid_master_value_is_not_truncated(self):
        with tempfile.TemporaryDirectory() as td:
            w=Workbook();s=w.active
            s.append(['标准图号','是否负责','是否启用','单箱数量'])
            s.append(['A','是','是',2.5]);w.save(Path(td)/'m.xlsx')
            self.assertEqual(load_product_master(Path(td)/'m.xlsx')['A'].units_per_carton,2.5)

    def test_country_logo_positive_and_prohibition_conflicts(self):
        req=extract_requirements(['合同号：ABC 车型：C1 国家：土耳其 LOGO类型：福田 包装模数：12 标签网址：https://example.com'])
        self.assertEqual(req.country,'土耳其')
        self.assertEqual(req.logo_type,'福田')
        self.assertEqual(req.errors,[])
        req=extract_requirements(['合同号：ABC 车型：C1 国家：土耳其 LOGO类型：福田',
                                '包装标签不得体现国家及福田Logo'])
        self.assertTrue(any('国家' in e and '冲突' in e for e in req.errors))
        self.assertTrue(any('LOGO' in e and '冲突' in e for e in req.errors))

    def test_delivery_time_keeps_dots_and_no_logo_maps_slash(self):
        req=extract_requirements(['送货时间：10.19-10.20日，包装模数：24. 国家：/ logo类型：无（选第二个）'])
        self.assertEqual(req.delivery_time,'10.19-10.20日')
        self.assertEqual(req.logo_type,'/')

    def test_auxiliary_sheet_is_ignored_but_plain_requirements_sheet_kept(self):
        with tempfile.TemporaryDirectory() as td:
            w=Workbook();s=w.active;s.append(['图号','数量']);s.append(['A',120])
            w.create_sheet('辅助').append(['-ck'])
            w.create_sheet('说明').append(['送货时间：10.19-10.20日，包装模数：24'])
            w.save(Path(td)/'a.xlsx')
            batch=load_batch(td)
            self.assertEqual(len(batch.rows),1)
            self.assertTrue(any('包装模数' in t for t in batch.requirement_texts))
