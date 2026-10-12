from __future__ import annotations
from dataclasses import dataclass, field
from pathlib import Path
from datetime import datetime, date
import re
import json
from typing import Any
from openpyxl import load_workbook
from .core import BatchRow, ProductMaster, ProductRecord, normalize_part_no

ALIASES = {
 "code":["物料编码","物料号","图号","产品图号","零件号","partno","partnumber","materialcode"],
 "cn":["物料名称","产品名称","中文名称","零件名称","名称"],
 "en":["英文名称","产品英文名称","englishname","english"],
 "quantity":["数量","计划数量","需求数量","qty","quantity"],
 "modulus":["模数","包装模数","包装数量","经济批量","每组数量","每批数量"],
 "units":["单箱数量","每箱数量","箱装数","装箱数量","一箱装多少","箱内数量"],
 "carton_size":["纸箱尺寸","箱尺寸","外箱尺寸","包装尺寸","纸箱规格"],
 "station":["工位号","工位","站位号","工位编号"],
 "weight":["重量说明","零部件重量","单件重量","重量"],
 "contract_no":["合同号","合同编号"],
 "model":["车型","车型代码"],
 "country":["国家","标签国家"],
 "logo_type":["logo类型","logo","品牌标识"],
 "label_url":["标签网址","标签打印网址"],
 "finish_date":["计划时间","完成日期","计划日期","生产完成日期"],
 "remark":["备注","要求","技术要求","相关要求","说明"],
}
def _key(value: Any) -> str:
 return re.sub(r"[\s_：:（）()\[\]{}\-+/]", "", str(value or "").strip().lower())
ALIAS_KEYS = {_key(a): f for f, xs in ALIASES.items() for a in xs}

@dataclass
class BatchData:
 rows: list[BatchRow] = field(default_factory=list)
 requirement_texts: list[str] = field(default_factory=list)
 files: list[str] = field(default_factory=list)
 warnings: list[str] = field(default_factory=list)
 requirement_sources: list[dict] = field(default_factory=list)
 @property
 def errors(self): return self.warnings

@dataclass
class ExportRequirements:
 contract_no: str="/"; model: str="/"; country: str="/"; logo_type: str="/"
 delivery_place: str="/"; delivery_time: str="/"; modulus: str="/"; label_url: str="/"
 general_requirements: str=""; source_text: str=""; errors: list[str]=field(default_factory=list)
def _cell(row,index): return row[index] if index is not None and index < len(row) else None
def _value(row,indexes,name):
 v=_cell(row,indexes.get(name)); return v if v is not None and str(v).strip() else None
def _header_indexes(values):
 return {ALIAS_KEYS[_key(v)]:i for i,v in enumerate(values) if _key(v) in ALIAS_KEYS}
def _looks_header(x): return "code" in x and "quantity" in x
def _row_text(values): return " ".join(str(v).strip() for v in values if v is not None and str(v).strip())

def load_batch(folder: str|Path, aliases_path=None)->BatchData:
 folder=Path(folder)
 if not folder.is_dir(): raise FileNotFoundError(f"批次文件夹不存在：{folder}")
 result=BatchData(); files=sorted(p for p in folder.iterdir() if p.suffix.lower() in {".xlsx",".xlsm",".xls"})
 if not files: raise ValueError("批次文件夹中没有 Excel 文件")
 alias_map = dict(ALIAS_KEYS)
 if aliases_path and Path(aliases_path).exists():
  extra = json.loads(Path(aliases_path).read_text(encoding='utf-8'))
  for name, aliases in extra.items():
   if name not in ALIASES: raise ValueError(f"未知字段别名键：{name}")
   for alias in aliases: alias_map[_key(alias)] = name
 for path in files:
  if path.name.startswith('~$'): continue
  if path.suffix.lower()==".xls":
   try:
    import xlrd
    old = xlrd.open_workbook(path)
    sheets = [(s.name, [[s.cell_value(r,c) for c in range(s.ncols)] for r in range(s.nrows)]) for s in old.sheets()]
   except Exception as exc:
    result.warnings.append(f"无法读取 {path.name}：{exc}"); continue
   book = None
  else:
   try:
    book=load_workbook(path,data_only=True,read_only=True)
    sheets=[(s.title,s.iter_rows(values_only=True)) for s in book.worksheets]
   except Exception as exc:
    result.warnings.append(f"无法读取 {path.name}：{exc}"); continue
  result.files.append(str(path))
  found = False
  for sheet_name, sheet_rows in sheets:
   indexes=None
   for number,values in enumerate(sheet_rows,1):
    values=tuple(values)
    if indexes is None:
     candidate={alias_map[_key(v)]:i for i,v in enumerate(values) if _key(v) in alias_map}
     if _looks_header(candidate): indexes=candidate; found=True; continue
     text=_row_text(values)
     if text and any(k in text for k in ("通知","合同","车型","网址","国家","包装","送货")):
      result.requirement_texts.append(text)
      result.requirement_sources.append({"file":path.name,"sheet":sheet_name,"row":number,"text":text})
     continue
    candidate={alias_map[_key(v)]:i for i,v in enumerate(values) if _key(v) in alias_map}
    if _looks_header(candidate): indexes=candidate; continue
    code=_value(values,indexes,"code"); quantity=_value(values,indexes,"quantity")
    if code is None and quantity is None:
     text=_row_text(values)
     if text:
      result.requirement_texts.append(text)
      result.requirement_sources.append({"file":path.name,"sheet":sheet_name,"row":number,"text":text})
     continue
    if code is None: continue
    fields={"英文名称":_value(values,indexes,"en"),"包装模数":_value(values,indexes,"modulus"),"单箱数量":_value(values,indexes,"units"),"纸箱尺寸":_value(values,indexes,"carton_size"),"工位号":_value(values,indexes,"station"),"重量说明":_value(values,indexes,"weight")}
    finish=_value(values,indexes,"finish_date")
    if finish:
     if isinstance(finish,(date,datetime)): finish=finish.date().isoformat() if isinstance(finish,datetime) else finish.isoformat()
     fields["完成日期"]=str(finish)
    for field_name,label in [("contract_no","合同号"),("model","车型"),("country","国家"),("logo_type","LOGO类型"),("label_url","标签网址")]:
     explicit=_value(values,indexes,field_name)
     if explicit is not None:
      fields[label]=str(explicit)
      result.requirement_texts.append(f"{label}：{explicit}")
      result.requirement_sources.append({"file":path.name,"sheet":sheet_name,"row":number,"text":f"{label}：{explicit}"})
    remark=_value(values,indexes,"remark")
    if remark:
     fields["要求"]=str(remark); result.requirement_texts.append(str(remark))
     result.requirement_sources.append({"file":path.name,"sheet":sheet_name,"row":number,"text":str(remark)})
     for label,alias in [("英文名称","en"),("包装模数","modulus"),("单箱数量","units"),("纸箱尺寸","carton_size"),("工位号","station"),("重量说明","weight")]:
      if not fields.get(label):
       keys="|".join(re.escape(a) for a in ALIASES[alias])
       match=re.search(r"(?:"+keys+r")[：:]\s*([^；;，,。\n]+)",str(remark),re.I)
       if match: fields[label]=match.group(1).strip()
    result.rows.append(BatchRow(str(code),str(_value(values,indexes,"cn") or ""),quantity,fields,path.name,sheet_name,number))
  if book: book.close()
  if not found: result.warnings.append(f"{path.name}：没有有效物料工作表，不能静默跳过")
 if not result.rows: raise ValueError("未识别到包含物料编码和数量的有效工作表")
 return result

def _match(text,patterns):
 for p in patterns:
  m=re.search(p,text,re.I)
  if m: return m.group(1).strip().rstrip("。；;，,")
 return "/"

def extract_requirements(texts)->ExportRequirements:
 source="\n".join(dict.fromkeys(t.strip() for t in texts if t and t.strip())).strip()
 req=ExportRequirements(source_text=source)
 req.contract_no=_match(source,[r"合同号[：:\s()（）]*([A-Z0-9-]+)",r"/(HW[A-Z0-9]+)"])
 req.model=_match(source,[r"车型[：:：\s()（）]*([A-Z0-9-]+)",r"/(C[A-Z0-9-]+)"])
 req.country=_match(source,[r"国家[：:：\s]*([^\s，,；;。]+)"])
 req.logo_type=_match(source,[r"logo类型[：:：\s（）()]*([^\s，,；;。]+)",r"LOGO类型[：:：\s（）()]*([^\s，,；;。]+)"])
 req.delivery_place=_match(source,[r"送货地点[：:：\s]*([^\n。；;，,]+)"])
 req.delivery_time=_match(source,[r"送货时间[：:：\s]*([^\n。；;，,]+)"])
 req.modulus=_match(source,[r"包装模数[：:：\s]*([0-9]+)",r"经济批量[为：:：\s]*([0-9]+)"])
 url=re.search(r"https?://[^\s，,；;）)]+",source,re.I); req.label_url=url.group(0).rstrip("。；;") if url else "/"

 country_values = re.findall(r"国家[：:]\s*([^\s，,；;。（）()]+)",source,re.I)
 logo_values = re.findall(r"logo(?:类型)?[：:]\s*([^\s，,；;。（）()]+)",source,re.I)
 def normalized_label(v):
  return "/" if v.strip().lower() in {"/","无","不带","不显示","none","no"} else v.strip()
 country_values = list(dict.fromkeys(normalized_label(v) for v in country_values))
 logo_values = list(dict.fromkeys(normalized_label(v) for v in logo_values))
 country_forbidden = bool(re.search(r"(?:不得体现|不得显示|不显示|不带|无)[^\n。；;，,]*国家",source,re.I))
 logo_forbidden = bool(re.search(r"(?:不得体现|不得显示|不显示|不带|无)[^\n。；;，,]*logo",source,re.I))
 if len(country_values)>1 or (country_forbidden and any(v!="/" for v in country_values)):
  req.errors.append("国家要求存在冲突")
 if len(logo_values)>1 or (logo_forbidden and any(v!="/" for v in logo_values)):
  req.errors.append("LOGO要求存在冲突")
 req.country = "/" if country_forbidden else (country_values[0] if country_values else "/")
 req.logo_type = "/" if logo_forbidden else (logo_values[0] if logo_values else "/")
 if req.contract_no=="/": req.errors.append("合同号缺失")
 if req.model=="/": req.errors.append("车型缺失")
 if req.label_url=="/": req.errors.append("标签网址缺失")
 contracts=re.findall(r"合同号[：:\s()（）]*([A-Z0-9-]+)",source,re.I)
 models=re.findall(r"车型[：:\s()（）]*([A-Z0-9-]+)",source,re.I)
 if len(set(contracts))>1: req.errors.append("合同号存在冲突："+ "、".join(dict.fromkeys(contracts)))
 if len(set(models))>1: req.errors.append("车型存在冲突："+ "、".join(dict.fromkeys(models)))

 if not country_values and not country_forbidden:
  req.errors.append("国家标签要求缺失")
 if not logo_values and not logo_forbidden:
  req.errors.append("LOGO标签要求缺失")
 # 保留原始要求供审核；正文从技术要求开始，排除群通知和重复网址。
 body=source
 body_start=re.search(r"(?:其他要求|技术要求)[：:]",body)
 if body_start: body=body[body_start.end():]
 body=re.sub(r"(?:\d+[）)]\s*)?(?:标签网址|标签打印网址)[：:]\s*https?://[^\s]+","",body)
 req.general_requirements=body.strip()
 return req

def load_product_master(path: str|Path)->ProductMaster:
 book=load_workbook(path,data_only=True,read_only=True); master=ProductMaster()
 aliases={"standard_code":["标准图号","图号","物料编码","物料号"],"responsible":["是否负责","负责"],"enabled":["是否启用","启用"],"chinese_name":["中文名称","产品名称"],"english_name":["英文名称","产品英文名称"],"packaging_modulus":["包装模数","模数","经济批量"],"units_per_carton":["单箱数量","每箱数量","箱装数"],"carton_size":["纸箱尺寸","箱尺寸"],"station_no":["工位号","工位"],"weight_note":["重量说明","重量"]}
 amap={_key(a):f for f,xs in aliases.items() for a in xs}
 for sheet in book.worksheets:
  rows=sheet.iter_rows(values_only=True); headers=next(rows,None)
  if not headers: continue
  indexes={amap[_key(v)]:i for i,v in enumerate(headers) if _key(v) in amap}
  if "standard_code" not in indexes: continue
  for values in rows:
   raw=_cell(values,indexes["standard_code"])
   if not raw: continue
   def val(name,default=""): return _cell(values,indexes.get(name)) if indexes.get(name) is not None else default
   def boolean(name,default=True): return str(val(name,default)).strip().lower() not in {"否","no","false","0"}
   def integer(name):
    value=val(name,None)
    return value if value not in (None,"") else None
   code=normalize_part_no(raw)
   if code in master:
    book.close()
    raise ValueError(f"产品资料表存在重复标准图号：{code}，请只保留一条记录")
   master[code]=ProductRecord(code,boolean("responsible"),boolean("enabled"),str(val("chinese_name","") or ""),str(val("english_name","") or ""),integer("packaging_modulus"),integer("units_per_carton"),str(val("carton_size","") or ""),str(val("station_no","") or ""),str(val("weight_note","") or ""))
 book.close()
 if not master: raise ValueError("产品资料表未识别到标准图号列")
 return master
