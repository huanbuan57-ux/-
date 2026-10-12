from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment


def write_report(analysis, output_path, extra_errors=()):
    workbook = Workbook()
    overview = workbook.active
    overview.title = '结果'
    overview.append(['项目', '内容'])
    overview.append(['生成状态', '阻断' if analysis.errors or extra_errors else '待人工审核'])
    overview.append(['负责产品数', len(analysis.rows)])
    overview.append(['排除产品数', len(analysis.excluded)])
    for error in dict.fromkeys([*analysis.errors, *extra_errors]):
        overview.append(['错误', error])
    for warning in analysis.warnings:
        overview.append(['提示', warning])
    reqs = workbook.create_sheet('要求')
    reqs.append(['字段', '最终采用值'])
    for key, label in [('contract_no','合同号'),('model','车型'),('country','国家'),('logo_type','LOGO类型'),
                       ('delivery_place','送货地点'),('delivery_time','送货时间'),('modulus','包装模数'),
                       ('label_url','标签网址'),('source_text','原始要求文本')]:
        reqs.append([label, getattr(analysis.requirements,key)])
    source_sheet=workbook.create_sheet('要求来源')
    source_sheet.append(['文件','工作表','行号','原始要求'])
    for item in analysis.requirement_sources:
        source_sheet.append([item['file'],item['sheet'],item['row'],item['text']])
    trace = workbook.create_sheet('字段来源')
    trace.append(['文件','工作表','行号','原始图号','标准图号','字段','Excel值','资料表值','最终采用值','来源/处理'])
    mapping={'中文名称':'chinese_name','英文名称':'english_name','包装模数':'packaging_modulus','单箱数量':'units_per_carton',
             '纸箱尺寸':'carton_size','工位号':'station_no','重量说明':'weight_note'}
    for row in analysis.all_rows:
        record=analysis.master.get(row.standard_code)
        if not row.resolved: continue
        for label, attr in mapping.items():
            raw=row.chinese_name if label=='中文名称' else row.fields.get(label)
            baseline=getattr(record,attr,None)
            source='Excel' if raw not in (None,'') else ('缺失按规则写 /' if label=='工位号' else '产品资料表')
            trace.append([row.source_file,row.source_sheet,row.source_row,row.raw_code,row.standard_code,label,
                          raw,baseline,row.resolved.get(label),source])
        for label in ['完成日期','要求','合同号','车型','国家','LOGO类型','标签网址']:
            if row.fields.get(label) is not None:
                trace.append([row.source_file,row.source_sheet,row.source_row,row.raw_code,row.standard_code,label,row.fields[label],None,row.fields[label],'Excel'])
        trace.append([row.source_file,row.source_sheet,row.source_row,row.raw_code,row.standard_code,'数量',
                      row.quantity,None,row.quantity,'Excel'])
        trace.append([row.source_file,row.source_sheet,row.source_row,row.raw_code,row.standard_code,'纸箱总数',
                      None,None,row.carton_count,'数量 ÷ 单箱数量'])
    excluded=workbook.create_sheet('排除清单')
    excluded.append(['文件','工作表','行号','原始图号','标准图号','原因'])
    for item in analysis.excluded:
        excluded.append([item[k] for k in ['source_file','source_sheet','source_row','raw_code','standard_code','reason']])
    for sheet in workbook:
        sheet.freeze_panes='A2'
        sheet.auto_filter.ref=sheet.dimensions
        for cell in sheet[1]:
            cell.font=Font(bold=True,color='FFFFFF')
            cell.fill=PatternFill('solid',fgColor='244A68')
        for cells in sheet:
            for cell in cells:
                cell.alignment=Alignment(vertical='top',wrap_text=True)
                # 将输入视作文字，避免字符串以 = 开头被导出为公式。
                if isinstance(cell.value,str): cell.data_type='s'
        for col in sheet.columns:
            sheet.column_dimensions[col[0].column_letter].width=min(65,max(14,max(len(str(c.value or '')) for c in col)+2))
    Path(output_path).parent.mkdir(parents=True,exist_ok=True)
    workbook.save(output_path)
    return Path(output_path)
